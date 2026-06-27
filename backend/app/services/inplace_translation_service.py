import json
import re
from pptx import Presentation

def _clean_json_response(text):
    """Clean markdown formatting from LLM JSON response."""
    import re
    match = re.search(r'\{.*\}', text, re.DOTALL)
    if match:
        return match.group(0)
    return text.strip()

def _translate_chunk(chunk_dict, target_language, instruction, enhancement_instructions, agent):
    """Translate a small chunk of strings."""
    from crewai import Task, Crew

    json_payload = json.dumps(chunk_dict, ensure_ascii=False, indent=2)

    translation_task = Task(
        description=f'''
        You are a professional Translator. Your primary task is to TRANSLATE the string values in the following JSON object into {target_language}.
        You MUST keep the exact same JSON keys. Only translate the values.

        TRANSLATION MODE INSTRUCTIONS:
        {instruction}

        {enhancement_instructions}

        CRITICAL RULES:
        1. Return ONLY a valid JSON object. No preamble, no explanation, no markdown blocks.
        2. Do NOT change the keys.
        3. Maintain any leading/trailing whitespace if possible.
        4. Keep numeric values, percentages, dates, and proper nouns untranslated if appropriate.

        JSON TO TRANSLATE:
        {json_payload}
        ''',
        expected_output="A valid JSON object containing the translated strings.",
        agent=agent
    )

    crew = Crew(
        agents=[agent],
        tasks=[translation_task],
        verbose=True
    )

    result = crew.kickoff()
    result_text = _clean_json_response(str(result))
    return json.loads(result_text)

def _batch_translate_dict(text_dict, target_language, instruction, enhancement_instructions, agent):
    """Translate a dictionary of strings using the LLM agent by splitting into chunks."""
    if not text_dict:
        return {}

    items = list(text_dict.items())
    chunk_size = 15
    translated_dict = {}

    for i in range(0, len(items), chunk_size):
        chunk = dict(items[i:i + chunk_size])
        print(f"Translating chunk {i//chunk_size + 1} of {(len(items) + chunk_size - 1)//chunk_size}...")
        try:
            chunk_result = _translate_chunk(chunk, target_language, instruction, enhancement_instructions, agent)
            for k, v in chunk_result.items():
                translated_dict[k] = v
        except Exception as e:
            print(f"Chunk translation failed: {e}")
            for k, v in chunk.items():
                translated_dict[k] = v

    # Fallback for any missing keys
    for k, v in text_dict.items():
        if k not in translated_dict:
            translated_dict[k] = v

    return translated_dict


def inplace_translate_pptx(source_path, output_path, target_language, instruction, enhancement_instructions, agent):
    """Extract text from PPTX, translate in batch, and re-inject in-place."""
    prs = Presentation(source_path)

    # 1. Extract
    text_map = {}
    counter = 0

    # Store references to runs so we can easily replace them later
    run_references = {}

    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        original_text = run.text
                        if original_text and original_text.strip():
                            key = str(counter)
                            text_map[key] = original_text
                            run_references[key] = run
                            counter += 1
            if shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        for paragraph in cell.text_frame.paragraphs:
                            for run in paragraph.runs:
                                original_text = run.text
                                if original_text and original_text.strip():
                                    key = str(counter)
                                    text_map[key] = original_text
                                    run_references[key] = run
                                    counter += 1

    # 2. Translate
    translated_map = _batch_translate_dict(
        text_map, target_language, instruction, enhancement_instructions, agent
    )

    # 3. Inject
    for key, run in run_references.items():
        if key in translated_map and translated_map[key]:
            run.text = translated_map[key]

    prs.save(output_path)
    return True


def inplace_translate_xlsx(source_path, output_path, target_language, instruction, enhancement_instructions, agent):
    """Extract text from XLSX, translate in batch, and re-inject in-place."""
    import openpyxl

    wb = openpyxl.load_workbook(source_path)

    # 1. Extract
    text_map = {}
    counter = 0
    cell_references = {}

    for sheet_name in wb.sheetnames:
        sheet = wb[sheet_name]
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value and isinstance(cell.value, str):
                    # Ignore formulas
                    if not cell.value.startswith('='):
                        original_text = cell.value
                        if original_text.strip():
                            key = str(counter)
                            text_map[key] = original_text
                            cell_references[key] = cell
                            counter += 1

    # 2. Translate
    translated_map = _batch_translate_dict(
        text_map, target_language, instruction, enhancement_instructions, agent
    )

    # 3. Inject
    for key, cell in cell_references.items():
        if key in translated_map and translated_map[key]:
            cell.value = translated_map[key]

    wb.save(output_path)
    return True

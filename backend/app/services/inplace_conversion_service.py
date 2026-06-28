import json
import re
import openpyxl
from pptx import Presentation
from docx import Document

def _clean_json_response(text):
    """Clean markdown formatting from LLM JSON response."""
    import re
    match = re.search(r'\{.*\}', text, re.DOTALL)
    if match:
        return match.group(0)
    return text.strip()

def _convert_chunk(chunk_dict, extracted_fields, agent):
    """Use the LLM to fill in a small chunk of template strings using the extracted fields."""
    from crewai import Task, Crew

    template_payload = json.dumps(chunk_dict, ensure_ascii=False, indent=2)
    fields_payload = json.dumps(extracted_fields, ensure_ascii=False, indent=2)

    conversion_task = Task(
        description=f'''
        You are a highly precise formatting assistant. Your task is to POPULATE the string values in a JSON template using the provided EXTRACTED FIELDS.
        
        The JSON template represents text chunks extracted from a formatted document template (like a resume, report, or presentation).
        Your job is to read the EXTRACTED FIELDS, find the relevant information, and REPLACE any placeholder text in the JSON template with the extracted information.

        CRITICAL RULES:
        1. Return ONLY a valid JSON object. No preamble, no explanation, no markdown blocks.
        2. Do NOT change the JSON keys. Keep them exactly as they are.
        3. If a JSON value contains a placeholder (e.g., "[Name]", "____", "Insert Date Here") and you have that data in EXTRACTED FIELDS, replace the placeholder with the fact.
        4. If a JSON value is a structural label (e.g., "Employee Name:", "CONFIDENTIAL"), append the extracted value if it's a combined label/placeholder (e.g., "Name: ____" becomes "Name: John Doe"). If it's just a label and the value belongs in a different chunk, leave the label untouched.
        5. If the required information is NOT found in the EXTRACTED FIELDS, leave the placeholder unchanged.
        6. Do NOT translate the text. Only fill in the blanks.
        7. If a chunk is just boilerplate text with no placeholders, return it EXACTLY as it is.

        EXTRACTED FIELDS:
        {fields_payload}

        JSON TEMPLATE TO POPULATE:
        {template_payload}
        ''',
        expected_output="A valid JSON object containing the populated template strings.",
        agent=agent
    )

    crew = Crew(
        agents=[agent],
        tasks=[conversion_task],
        verbose=True
    )

    result = crew.kickoff()
    result_text = _clean_json_response(str(result))
    return json.loads(result_text)

def _batch_convert_dict(text_dict, extracted_fields, agent):
    """Convert a dictionary of strings using the LLM agent by splitting into chunks."""
    if not text_dict:
        return {}
    if not extracted_fields:
        return text_dict

    items = list(text_dict.items())
    chunk_size = 15
    converted_dict = {}

    for i in range(0, len(items), chunk_size):
        chunk = dict(items[i:i + chunk_size])
        print(f"Populating chunk {i//chunk_size + 1} of {(len(items) + chunk_size - 1)//chunk_size}...")
        try:
            chunk_result = _convert_chunk(chunk, extracted_fields, agent)
            for k, v in chunk_result.items():
                converted_dict[k] = v
        except Exception as e:
            print(f"Chunk conversion failed: {e}")
            for k, v in chunk.items():
                converted_dict[k] = v

    # Fallback for any missing keys
    for k, v in text_dict.items():
        if k not in converted_dict:
            converted_dict[k] = v

    return converted_dict


def inplace_convert_docx(reference_path, output_path, data, agent):
    """Extract text from DOCX template, fill blanks via LLM, and re-inject in-place."""
    doc = Document(reference_path)
    text_map = {}
    counter = 0
    run_references = {}

    # Extract
    for p in doc.paragraphs:
        for run in p.runs:
            if run.text and run.text.strip():
                key = str(counter)
                text_map[key] = run.text
                run_references[key] = run
                counter += 1
                
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for run in p.runs:
                        if run.text and run.text.strip():
                            key = str(counter)
                            text_map[key] = run.text
                            run_references[key] = run
                            counter += 1

    # Populate
    fields = data.get("fields", {})
    converted_map = _batch_convert_dict(text_map, fields, agent)

    # Inject
    for key, run in run_references.items():
        if key in converted_map and converted_map[key]:
            run.text = converted_map[key]

    # --- PHASE 2: TABULAR ROW APPENDING ---
    tables_data = data.get("tables", [])
    
    for t in doc.tables:
        if len(t.rows) > 0:
            doc_headers = [c.text.strip().lower() for c in t.rows[0].cells if c.text.strip()]
            if not doc_headers:
                continue
                
            matched_tbl_data = None
            for tbl_data in tables_data:
                tbl_headers = tbl_data.get("headers", [])
                match_count = sum(1 for h in tbl_headers if h and str(h).strip().lower() in doc_headers)
                if match_count >= 2 or (len(doc_headers) == 1 and match_count == 1):
                    matched_tbl_data = tbl_data
                    break
                    
            rows_to_add = []
            if matched_tbl_data:
                tbl_headers = matched_tbl_data.get("headers", [])
                for row_data in matched_tbl_data.get("rows", []):
                    row_dict = {}
                    for i, val in enumerate(row_data):
                        if i < len(tbl_headers) and tbl_headers[i]:
                            row_dict[str(tbl_headers[i]).lower()] = val
                    rows_to_add.append(row_dict)
            else:
                match_count = sum(1 for k in fields.keys() if str(k).lower() in doc_headers)
                if match_count >= 2 or (len(doc_headers) == 1 and match_count == 1):
                    row_dict = {str(k).lower(): v for k, v in fields.items()}
                    rows_to_add.append(row_dict)
            
            if rows_to_add:
                doc_headers_raw = [c.text.strip() for c in t.rows[0].cells]
                for row_dict in rows_to_add:
                    row_cells = t.add_row().cells
                    for i, doc_h in enumerate(doc_headers_raw):
                        if i < len(row_cells) and doc_h:
                            matched_val = ""
                            for k, v in row_dict.items():
                                if k in doc_h.lower() or doc_h.lower() in k:
                                    if isinstance(v, dict) and "value" in v:
                                        matched_val = v["value"]
                                    else:
                                        matched_val = v
                                    break
                            row_cells[i].text = str(matched_val)

    doc.save(output_path)
    return True


def inplace_convert_pptx(reference_path, output_path, extracted_fields, agent):
    """Extract text from PPTX template, fill blanks via LLM, and re-inject in-place."""
    prs = Presentation(reference_path)
    text_map = {}
    counter = 0
    run_references = {}

    # Extract
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    for run in p.runs:
                        if run.text and run.text.strip():
                            key = str(counter)
                            text_map[key] = run.text
                            run_references[key] = run
                            counter += 1
            if shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        for p in cell.text_frame.paragraphs:
                            for run in p.runs:
                                if run.text and run.text.strip():
                                    key = str(counter)
                                    text_map[key] = run.text
                                    run_references[key] = run
                                    counter += 1

    # Populate
    converted_map = _batch_convert_dict(text_map, extracted_fields, agent)

    # Inject
    for key, run in run_references.items():
        if key in converted_map and converted_map[key]:
            run.text = converted_map[key]

    prs.save(output_path)
    return True


def inplace_convert_xlsx(reference_path, output_path, data, agent=None):
    """Support both Form-based in-place replacement and Tabular data appending."""
    import openpyxl
    wb = openpyxl.load_workbook(reference_path)
    ws = wb.active

    # Find headers (assuming row 1)
    headers = []
    for cell in ws[1]:
        if cell.value:
            headers.append(str(cell.value).strip())

    is_pure_tabular = ws.max_row <= 2 and len(headers) > 1

    # --- PHASE 1: FORM-BASED IN-PLACE REPLACEMENT ---
    # Only run if it's a form (not just a pure table header)
    if not is_pure_tabular:
        text_map = {}
        counter = 0
        cell_references = {}
        
        for sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.value and isinstance(cell.value, str):
                        if not cell.value.startswith('='):
                            if cell.value.strip():
                                key = str(counter)
                                text_map[key] = cell.value
                                cell_references[key] = cell
                                counter += 1

        if text_map and agent:
            fields = data.get("fields", {})
            converted_map = _batch_convert_dict(text_map, fields, agent)
            
            for key, cell in cell_references.items():
                if key in converted_map and converted_map[key]:
                    cell.value = converted_map[key]

    # --- PHASE 2: TABULAR ROW APPENDING ---
    tables = data.get("tables", [])
    fields = data.get("fields", {})
    
    rows_to_write = []
    if tables:
        target_table = None
        for tbl in tables:
            tbl_headers = tbl.get("headers", [])
            # Require at least 2 matching headers to confidently match the table
            match_count = sum(1 for h in tbl_headers if h and h.lower() in [c.lower() for c in headers])
            if match_count >= 2 or (len(headers) == 1 and match_count == 1):
                target_table = tbl
                break
                
        if target_table:
            tbl_headers = target_table.get("headers", [])
            for row_data in target_table.get("rows", []):
                row_dict = dict(fields)
                for i, val in enumerate(row_data):
                    if i < len(tbl_headers) and tbl_headers[i]:
                        row_dict[tbl_headers[i]] = val
                rows_to_write.append(row_dict)
                
    if not rows_to_write and is_pure_tabular:
        rows_to_write.append(fields)

    if rows_to_write:
        current_row = ws.max_row + 1
        for row_dict in rows_to_write:
            for col_idx, col_name in enumerate(headers, start=1):
                matched_val = ""
                for k, v in row_dict.items():
                    if k.lower() in col_name.lower() or col_name.lower() in k.lower():
                        if isinstance(v, dict) and "value" in v:
                            matched_val = str(v["value"])
                        else:
                            matched_val = str(v)
                        break
                
                ws.cell(row=current_row, column=col_idx, value=matched_val)
            current_row += 1

    wb.save(output_path)
    return True

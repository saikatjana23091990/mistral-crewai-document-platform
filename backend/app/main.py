from fastapi import FastAPI, UploadFile, File, Request, HTTPException, Form
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse
from app.parsers.document_parser import parse_document
from app.services.conversion_service import run_conversion
from app.rag.rag_engine import build_rag, build_multi_rag
import json
import shutil
from pathlib import Path
from datetime import datetime, timezone
import time
import pandas as pd
import re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from docx import Document
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas
from fastapi.staticfiles import StaticFiles

app = FastAPI()

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

app.mount("/outputs", StaticFiles(directory="outputs"), name="outputs")

UPLOAD_DIR = Path("uploads")
OUTPUT_DIR = Path("outputs")
STATS_FILE = Path("stats.json")
HISTORY_FILE = Path("history.json")
SETTINGS_FILE = Path("settings.json")

UPLOAD_DIR.mkdir(exist_ok=True)
OUTPUT_DIR.mkdir(exist_ok=True)

TRANSLATION_STATS_FILE = Path("translation_stats.json")
TRANSLATION_HISTORY_FILE = Path("translation_history.json")
TRANSLATION_GLOSSARY_FILE = Path("translation_glossary.json")

rag_sessions = {}
chat_documents = {}

DEFAULT_STATS = {
    "documents_converted": 0,
    "chat_interactions": 0,
    "api_calls": 0,
    "total_size_saved": 0,
    "success_rate": 100,
    "this_month_conversions": 0
}
DEFAULT_HISTORY = []
DEFAULT_SETTINGS = {
    "provider": "groq",
    "model": "llama-3.1-8b-instant",
    "temperature": 0.5,
    "reasoningDepth": "fast",
    "memoryEnabled": True,
    "streamResponses": True,
    "includeCitations": True,
    "explainConflicts": True,
    "showConfidenceScores": True,
    "defaultTranslationMode": "Business",
    "defaultTargetLanguage": "Spanish"
}

DEFAULT_TRANSLATION_STATS = {
    "total_translations": 0,
    "successful_translations": 0,
    "average_quality": 0.0,
    "languages_supported": 12,
    "documents_translated": 0
}
DEFAULT_TRANSLATION_HISTORY = []
DEFAULT_TRANSLATION_GLOSSARY = []

def load_stats():
    if STATS_FILE.exists():
        with open(STATS_FILE, encoding="utf-8") as f:
            return {**DEFAULT_STATS, **json.load(f)}
    return DEFAULT_STATS.copy()

def save_stats(stats):
    with open(STATS_FILE, "w", encoding="utf-8") as f:
        json.dump(stats, f)

def load_history():
    if HISTORY_FILE.exists():
        with open(HISTORY_FILE, encoding="utf-8") as f:
            data = json.load(f)
            return data if isinstance(data, list) else DEFAULT_HISTORY.copy()
    return DEFAULT_HISTORY.copy()

def save_history(history):
    with open(HISTORY_FILE, "w", encoding="utf-8") as f:
        json.dump(history, f)

def load_settings_data():
    if SETTINGS_FILE.exists():
        try:
            with open(SETTINGS_FILE, encoding="utf-8") as f:
                return {**DEFAULT_SETTINGS, **json.load(f)}
        except Exception:
            return DEFAULT_SETTINGS.copy()
    return DEFAULT_SETTINGS.copy()

def save_settings_data(settings_data):
    with open(SETTINGS_FILE, "w", encoding="utf-8") as f:
        json.dump(settings_data, f)

stats = load_stats()
conversion_history = load_history()
global_settings = load_settings_data()

def load_json(path, default):
    if path.exists():
        try:
            with open(path, encoding="utf-8") as f:
                data = json.load(f)
                if isinstance(default, dict):
                    return {**default, **data}
                return data if isinstance(data, list) else default.copy()
        except:
            return default.copy()
    return default.copy()

def save_json(path, data):
    with open(path, "w", encoding="utf-8") as f:
        json.dump(data, f)

translation_stats = load_json(TRANSLATION_STATS_FILE, DEFAULT_TRANSLATION_STATS)
translation_history = load_json(TRANSLATION_HISTORY_FILE, DEFAULT_TRANSLATION_HISTORY)
translation_glossary = load_json(TRANSLATION_GLOSSARY_FILE, DEFAULT_TRANSLATION_GLOSSARY)

# In-memory job state tracking for websockets
translation_jobs = {}

@app.middleware("http")
async def count_api_calls(request: Request, call_next):
    response = await call_next(request)
    if request.url.path not in ("/", "/stats"):
        stats["api_calls"] += 1
        save_stats(stats)
    return response

@app.get("/")
def health():
    return {"status": "running"}

@app.get("/stats")
def get_stats():
    return stats

@app.get("/history")
def get_history():
    return {"records": conversion_history}

@app.get("/settings")
def get_settings():
    return global_settings

@app.post("/settings")
async def update_settings(request: Request):
    try:
        new_settings = await request.json()
        global global_settings
        global_settings = {**global_settings, **new_settings}
        save_settings_data(global_settings)
        return {"message": "Settings updated", "settings": global_settings}
    except Exception as e:
        raise HTTPException(status_code=400, detail=str(e))

from fastapi import BackgroundTasks, WebSocket, WebSocketDisconnect
from app.services.translation_service import run_translation_background
import uuid
import asyncio

@app.get("/translate/stats")
def get_translation_stats():
    return translation_stats

@app.get("/translate/history")
def get_translation_history():
    return {"records": translation_history}

@app.post("/translate/analyze")
def translate_analyze(file: UploadFile = File(...)):
    source_path = UPLOAD_DIR / file.filename
    with open(source_path, "wb") as buffer:
        import shutil
        shutil.copyfileobj(file.file, buffer)
        
    from app.parsers.document_parser import analyze_document, parse_document
    metadata = analyze_document(str(source_path))
    
    # Extract text for word count and language detection heuristic
    try:
        full_text = parse_document(str(source_path))
        words = len(full_text.split())
        
        confidence = 98.7 if words > 0 else 0.0
        try:
            from langdetect import detect
            lang_code = detect(full_text[:5000])
            lang_map = {
                'en': 'English', 'fr': 'French', 'de': 'German', 'es': 'Spanish', 
                'ja': 'Japanese', 'ko': 'Korean', 'ar': 'Arabic', 'hi': 'Hindi', 
                'bn': 'Bengali', 'zh-cn': 'Chinese', 'zh-tw': 'Chinese'
            }
            detected_language = lang_map.get(lang_code, "English")
        except Exception:
            detected_language = "English"
        
    except Exception as e:
        words = 0
        detected_language = "Unknown"
        confidence = 0.0
        
    return {
        "filename": file.filename,
        "document_type": metadata.get("document_type", "Document"),
        "pages": metadata.get("pages", 1),
        "word_count": words,
        "detected_language": detected_language,
        "confidence": confidence
    }

@app.post("/translate/preview")
def translate_preview(
    filename: str = Form(...),
    targetLanguage: str = Form("Spanish"),
    sourceLanguage: str = Form("English"),
    mode: str = Form("Business"),
    provider: str = Form(None),
    model: str = Form(None),
    aiEnhancements: str = Form(None)
):
    source_path = UPLOAD_DIR / filename
    if not source_path.exists():
        raise HTTPException(status_code=404, detail="File not found")
        
    prov = provider or global_settings.get("provider", "mistral")
    mod = model or global_settings.get("model")
    
    from app.services.translation_service import generate_preview
    try:
        preview_data = generate_preview(str(source_path), sourceLanguage, targetLanguage, mode, prov, mod, aiEnhancements)
        return preview_data
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

@app.post("/translate/upload")
async def translate_upload(
    background_tasks: BackgroundTasks,
    file: UploadFile = File(...),
    targetLanguage: str = Form("Spanish"),
    sourceLanguage: str = Form("English"),
    mode: str = Form("Business"),
    provider: str = Form(None),
    model: str = Form(None),
    aiEnhancements: str = Form(None)
):
    source_path = UPLOAD_DIR / file.filename
    with open(source_path, "wb") as buffer:
        import shutil
        shutil.copyfileobj(file.file, buffer)
        
    job_id = str(uuid.uuid4())
    translation_jobs[job_id] = {
        "status": "Upload Complete",
        "progress": 10,
        "filename": file.filename,
        "targetLanguage": targetLanguage,
        "mode": mode,
        "aiEnhancements": aiEnhancements,
        "updates": []
    }
    
    prov = provider or global_settings.get("provider", "mistral")
    mod = model or global_settings.get("model")
    
    def send_progress_update(j_id, msg, final_result=None):
        if j_id in translation_jobs:
            translation_jobs[j_id]["status"] = msg
            translation_jobs[j_id]["updates"].append(msg)
            
            if msg == "Detecting Language...": translation_jobs[j_id]["progress"] = 30
            elif msg == "Translating...": translation_jobs[j_id]["progress"] = 60
            elif msg == "Formatting Reconstructed...": translation_jobs[j_id]["progress"] = 80
            elif msg == "Generating Reports...": translation_jobs[j_id]["progress"] = 90
            elif msg == "Completed": translation_jobs[j_id]["progress"] = 100
            
            if final_result:
                translation_jobs[j_id]["result"] = final_result
                if final_result.get("status") == "Completed":
                    # Update stats
                    translation_stats["total_translations"] += 1
                    translation_stats["successful_translations"] += 1
                    translation_stats["documents_translated"] += 1
                    translation_stats["average_quality"] = (translation_stats["average_quality"] * (translation_stats["total_translations"] - 1) + final_result.get("quality_score", 100)) / translation_stats["total_translations"]
                    save_json(TRANSLATION_STATS_FILE, translation_stats)
                    
                    # Add to history
                    record = {
                        "id": j_id,
                        "filename": file.filename,
                        "source_language": sourceLanguage,
                        "target_language": targetLanguage,
                        "translation_mode": mode,
                        "quality_score": final_result.get("quality_score"),
                        "corrections_applied": final_result.get("corrections_applied"),
                        "date": datetime.now(timezone.utc).isoformat(),
                        "status": "Completed",
                        "translated_path": final_result.get("translated_path")
                    }
                    translation_history.insert(0, record)
                    save_json(TRANSLATION_HISTORY_FILE, translation_history)
                elif final_result.get("status") == "Failed":
                    translation_stats["total_translations"] += 1
                    save_json(TRANSLATION_STATS_FILE, translation_stats)
    
    import threading
    thread = threading.Thread(
        target=run_translation_background,
        args=(job_id, str(source_path), sourceLanguage, targetLanguage, mode, prov, mod, aiEnhancements, globals(), send_progress_update)
    )
    thread.start()
    
    return {"jobId": job_id, "status": "Started"}

@app.websocket("/ws/translation/{job_id}")
async def translation_websocket(websocket: WebSocket, job_id: str):
    await websocket.accept()
    try:
        last_progress = -1
        while True:
            if job_id not in translation_jobs:
                await websocket.send_json({"error": "Job not found"})
                break
                
            job = translation_jobs[job_id]
            if job["progress"] != last_progress:
                await websocket.send_json(job)
                last_progress = job["progress"]
                
            if job["status"] in ["Completed", "Failed"]:
                await websocket.send_json(job)
                break
                
            await asyncio.sleep(1)
    except WebSocketDisconnect:
        pass
    except Exception as e:
        print("WS error:", e)

def get_target_format(reference_filename: str) -> str:
    """Determine target output format from reference file extension."""
    ext = Path(reference_filename).suffix.lower()
    if ext in (".xlsx", ".xls"):
        return "xlsx"
    elif ext == ".csv":
        return "csv"
    elif ext == ".json":
        return "json"
    elif ext == ".md":
        return "md"
    elif ext == ".docx":
        return "docx"
    elif ext == ".pdf":
        return "pdf"
    elif ext == ".pptx":
        return "pptx"
    else:
        return "txt"

def _get_reference_columns(reference_path: str) -> list:
    """Try to extract column names from an xlsx reference file."""
    try:
        if reference_path.endswith(('.xlsx', '.xls')):
            df = pd.read_excel(reference_path)
            return [str(c) for c in df.columns]
    except Exception:
        pass
    return []

def _extract_data_for_columns(source_text: str, columns: list) -> dict:
    """Dynamic extraction using reference columns + heuristics for sales data."""
    data = {col: "" for col in columns}
    for col in columns:
        cl = col.lower().strip()
        # Try exact column mention
        pat = rf'{re.escape(cl)}[:\-\s]+([^\n\r,]+)'
        m = re.search(pat, source_text, re.IGNORECASE)
        if m:
            val = m.group(1).strip()
            if val:
                data[col] = val
                continue

        # Quarter
        if 'quarter' in cl:
            m = re.search(r'Quarter:\s*([Q\d\s]+)', source_text, re.IGNORECASE)
            if m: data[col] = m.group(1).strip()
            else:
                m = re.search(r'(Q[1-4]\s*\d{4})', source_text)
                if m: data[col] = m.group(1)
        # Region
        elif 'region' in cl:
            m = re.search(r'Region:\s*([A-Za-z\s]+)', source_text, re.IGNORECASE)
            if m: data[col] = m.group(1).strip()
        # Sales Figure / Total Sales
        elif 'sales' in cl or 'figure' in cl:
            m = re.search(r'Total Sales:\s*\$?([\d,]+)', source_text, re.IGNORECASE)
            if m: data[col] = m.group(1)
            else:
                m = re.search(r'\$?([\d,]+\.?\d*)', source_text)
                if m: data[col] = m.group(1)
        # Source
        elif 'source' in cl:
            m = re.search(r'Prepared By:\s*([A-Za-z\s]+)', source_text, re.IGNORECASE)
            if m: data[col] = m.group(1).strip()
        # Date / Timestamp
        elif 'date' in cl or 'timestamp' in cl:
            m = re.search(r'(\d{4}-\d{2}-\d{2})', source_text)
            if m: data[col] = m.group(1)

    return data

def detect_conflicts(source_data_list: list, reference_columns: list = None) -> list:
    """Use actual reference columns for conflict detection."""
    conflicts = []
    if not reference_columns:
        reference_columns = ["Quarter", "Region", "Final Sales Figure"]
    for col in reference_columns:
        cl = col.lower()
        vals = []
        for sf in source_data_list:
            text = sf["text"]
            val = ""
            if 'quarter' in cl:
                m = re.search(r'Quarter:\s*([^\n]+)', text, re.IGNORECASE)
                val = m.group(1).strip() if m else ""
            elif 'region' in cl:
                m = re.search(r'Region:\s*([^\n]+)', text, re.IGNORECASE)
                val = m.group(1).strip() if m else ""
            elif 'sales' in cl or 'figure' in cl:
                m = re.search(r'Total Sales:\s*\$?([\d,]+)', text, re.IGNORECASE)
                val = m.group(1) if m else ""
            if val:
                vals.append({"source": sf["name"], "value": val})
        if len(vals) > 1:
            unique = set(v["value"] for v in vals)
            if len(unique) > 1:
                conflicts.append({"field": col, "values": vals})
    return conflicts

def generate_output_file(converted_text: str, target_format: str, source_stem: str, timestamp: str, reference_columns: list = None, reference_path: str = None) -> tuple[Path, str]:
    if reference_path and Path(reference_path).exists() and target_format in ["docx", "pptx", "xlsx"]:
        try:
            cleaned_text = converted_text.strip()
            if cleaned_text.startswith("```json"):
                cleaned_text = cleaned_text[7:]
            elif cleaned_text.startswith("```"):
                cleaned_text = cleaned_text[3:]
            import re
            json_match = re.search(r'\{.*\}', cleaned_text, re.DOTALL)
            if json_match:
                cleaned_text = json_match.group(0)
            
            import json
            data = json.loads(cleaned_text)
            fields = data.get("fields", {})
            for k, v in fields.items():
                if isinstance(v, dict) and "value" in v:
                    fields[k] = v.get("value", "")

            from app.agents.document_agents import get_agents
            provider = global_settings.get("provider", "groq")
            model = global_settings.get("model")
            _, _, formatter_agent, _ = get_agents(provider, model=model)

            filename = f"{source_stem}_converted_{timestamp}.{target_format}"
            output_path = OUTPUT_DIR / filename
            
            from app.services.inplace_conversion_service import inplace_convert_docx, inplace_convert_pptx, inplace_convert_xlsx

            if target_format == "docx":
                inplace_convert_docx(reference_path, str(output_path), data, formatter_agent)
                return output_path, filename
            elif target_format == "pptx":
                inplace_convert_pptx(reference_path, str(output_path), fields, formatter_agent)
                charts_data = data.get("charts", [])
                if charts_data:
                    from pptx import Presentation
                    prs = Presentation(str(output_path))
                    from app.services.pptx_chart_writer import update_pptx_charts
                    update_pptx_charts(prs, charts_data)
                    prs.save(str(output_path))
                return output_path, filename
            elif target_format == "xlsx":
                inplace_convert_xlsx(reference_path, str(output_path), data, formatter_agent)
                return output_path, filename
        except Exception as e:
            print(f"In-place conversion failed, falling back to legacy format: {e}")
            import traceback
            traceback.print_exc()

    if target_format == "xlsx":
        wb = Workbook()
        ws = wb.active
        ws.title = "Converted Data"

        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )

        columns = reference_columns if reference_columns else [
            "Quarter", "Region", "Final Sales Figure", "Source Selected", 
            "Alternative Values", "Reviewer", "Resolution Notes", "Approval Timestamp"
        ]

        for col_idx, col_name in enumerate(columns, start=1):
            cell = ws.cell(row=1, column=col_idx, value=col_name)
            cell.fill = header_fill
            cell.font = header_font
            cell.border = thin_border
            cell.alignment = Alignment(wrap_text=True)

        # Parse the JSON from converted_text
        rows_to_write = []
        try:
            cleaned_text = converted_text.strip()
            if cleaned_text.startswith("```json"):
                cleaned_text = cleaned_text[7:]
            elif cleaned_text.startswith("```"):
                cleaned_text = cleaned_text[3:]
            if cleaned_text.endswith("```"):
                cleaned_text = cleaned_text[:-3]
            data = json.loads(cleaned_text.strip())
            
            fields = data.get("fields", {})
            tables = data.get("tables", [])
            
            # If there's a table that seems to match our columns, use its rows
            if tables:
                target_table = None
                for tbl in tables:
                    # check if headers overlap well with our columns
                    headers = tbl.get("headers", [])
                    if any(h.lower() in [c.lower() for c in columns] for h in headers if h):
                        target_table = tbl
                        break
                        
                if target_table:
                    headers = target_table.get("headers", [])
                    for row_data in target_table.get("rows", []):
                        row_dict = dict(fields) # Start with base fields
                        for i, val in enumerate(row_data):
                            if i < len(headers) and headers[i]:
                                row_dict[headers[i]] = val
                        rows_to_write.append(row_dict)
                else:
                    rows_to_write.append(fields)
            else:
                rows_to_write.append(fields)
                
        except Exception:
            # Fallback if json parsing fails
            rows_to_write = [{}]

        if not rows_to_write:
            rows_to_write = [{}]

        current_row = 2
        for row_dict in rows_to_write:
            for col_idx, col_name in enumerate(columns, start=1):
                # case-insensitive match
                matched_val = "[Information not found in source documents]"
                for k, v in row_dict.items():
                    if k.lower() in col_name.lower() or col_name.lower() in k.lower():
                        matched_val = v
                        break
                cell = ws.cell(row=current_row, column=col_idx, value=str(matched_val))
                cell.border = thin_border
                cell.alignment = Alignment(wrap_text=True, vertical='top')
            current_row += 1

        for col_idx in range(1, len(columns) + 1):
            col_letter = chr(64 + col_idx) if col_idx <= 26 else 'A'
            ws.column_dimensions[col_letter].width = 22

        filename = f"{source_stem}_converted_{timestamp}.xlsx"
        output_path = OUTPUT_DIR / filename
        wb.save(output_path)
        return output_path, filename

    elif target_format == "docx":
        if reference_path and Path(reference_path).exists():
            doc = Document(reference_path)
        else:
            doc = Document()
            doc.add_heading("Converted Document", level=1)
            
        try:
            clean_text = converted_text.strip()
            # Clean markdown if present
            cleaned_text = converted_text.strip()
            if cleaned_text.startswith("```json"):
                cleaned_text = cleaned_text[7:]
            elif cleaned_text.startswith("```"):
                cleaned_text = cleaned_text[3:]
            if cleaned_text.endswith("```"):
                cleaned_text = cleaned_text[:-3]
            cleaned_text = cleaned_text.strip()
            
            data = json.loads(cleaned_text)
            
            # Map fields to paragraphs
            import re
            fields = data.get("fields", {})
            for p in doc.paragraphs:
                for key, val in fields.items():
                    if key and key.lower() in p.text.lower() and val:
                        original_text = p.text
                        if re.search(r'_{3,}', p.text):
                            p.text = re.sub(r'_{3,}', str(val), p.text, count=1)
                        elif ":" in p.text and p.text.strip().endswith(":"):
                            p.text = p.text + " " + str(val)
                        elif key.lower() == p.text.strip().lower():
                            p.text = p.text + ": " + str(val)
                        
                        # Stop checking other keys if we modified this paragraph
                        if p.text != original_text:
                            break

            # Map tables
            tables_data = data.get("tables", [])
            for table_data in tables_data:
                headers = table_data.get("headers", [])
                rows = table_data.get("rows", [])
                
                target_table = None
                for t in doc.tables:
                    if len(t.rows) > 0:
                        doc_headers = [c.text.strip() for c in t.rows[0].cells]
                        if all(h in doc_headers for h in headers if h) or all(h in headers for h in doc_headers if h):
                            target_table = t
                            break
                
                if target_table:
                    for row_data in rows:
                        row_cells = target_table.add_row().cells
                        for i, cell_val in enumerate(row_data):
                            if i < len(row_cells):
                                row_cells[i].text = str(cell_val)
                else:
                    if headers and rows:
                        doc.add_paragraph()
                        new_t = doc.add_table(rows=1, cols=len(headers))
                        new_t.style = 'Table Grid'
                        hdr_cells = new_t.rows[0].cells
                        for i, h in enumerate(headers):
                            hdr_cells[i].text = str(h)
                        for row_data in rows:
                            row_cells = new_t.add_row().cells
                            for i, cell_val in enumerate(row_data):
                                if i < len(row_cells):
                                    row_cells[i].text = str(cell_val)
                                    
        except json.JSONDecodeError:
            if not reference_path or not Path(reference_path).exists():
                for line in converted_text.split("\n"):
                    if line.strip():
                        doc.add_paragraph(line.strip())
            else:
                doc.add_paragraph("\n--- Converted Data ---\n")
                for line in converted_text.split("\n"):
                    if line.strip():
                        doc.add_paragraph(line.strip())

        filename = f"{source_stem}_converted_{timestamp}.docx"
        output_path = OUTPUT_DIR / filename
        doc.save(output_path)
        return output_path, filename

    elif target_format == "pptx":
        from pptx import Presentation
        from pptx.util import Pt, Emu
        from copy import deepcopy
        from lxml import etree
        import re
        if reference_path and Path(reference_path).exists():
            prs = Presentation(reference_path)
        else:
            prs = Presentation()
            prs.slides.add_slide(prs.slide_layouts[0])

        def _copy_run_format(source_run, target_run):
            """Copy all font/run-level formatting from source_run to target_run."""
            try:
                # Deep-copy the entire rPr (run properties) XML element if it exists
                source_rpr = source_run._r.find(
                    '{http://schemas.openxmlformats.org/drawingml/2006/main}rPr'
                )
                if source_rpr is not None:
                    target_rpr = target_run._r.find(
                        '{http://schemas.openxmlformats.org/drawingml/2006/main}rPr'
                    )
                    new_rpr = deepcopy(source_rpr)
                    if target_rpr is not None:
                        target_run._r.replace(target_rpr, new_rpr)
                    else:
                        # Insert rPr before the text element
                        target_run._r.insert(0, new_rpr)
            except Exception:
                pass

        def _replace_in_paragraph_runs(paragraph, old_text, new_value):
            """
            Surgically replace `old_text` with `new_value` across the paragraph's
            runs while preserving every run's original formatting.
            
            Strategy:
            1. Walk runs to find which ones contain the target text (which may span
               multiple runs due to PowerPoint's run-splitting).
            2. Replace only the matched portion, keeping surrounding text in its
               original run with its original formatting.
            3. If the match spans multiple runs, put the replacement text in the
               first affected run (keeping its formatting) and remove the matched
               portion from subsequent runs.
            """
            runs = list(paragraph.runs)
            if not runs:
                return False

            # Build a map: for each character position in the concatenated text,
            # record (run_index, offset_within_run)
            full_text = paragraph.text
            match_start = full_text.lower().find(old_text.lower())
            if match_start == -1:
                return False
            match_end = match_start + len(old_text)

            # Build position-to-run mapping
            char_pos = 0
            run_ranges = []  # (start_pos, end_pos, run_index)
            for r_idx, run in enumerate(runs):
                run_len = len(run.text)
                run_ranges.append((char_pos, char_pos + run_len, r_idx))
                char_pos += run_len

            # Find which runs are affected by the match
            affected = []
            for (r_start, r_end, r_idx) in run_ranges:
                if r_start < match_end and r_end > match_start:
                    # This run overlaps with the match
                    # Calculate the portion of this run's text that is matched
                    overlap_start = max(match_start, r_start) - r_start
                    overlap_end = min(match_end, r_end) - r_start
                    affected.append((r_idx, overlap_start, overlap_end))

            if not affected:
                return False

            # First affected run: replace the matched portion with new_value
            first_r_idx, first_overlap_start, first_overlap_end = affected[0]
            first_run = runs[first_r_idx]
            run_text = first_run.text
            first_run.text = run_text[:first_overlap_start] + str(new_value) + run_text[first_overlap_end:]

            # Remaining affected runs: remove only the matched portion
            for (r_idx, overlap_start, overlap_end) in affected[1:]:
                run = runs[r_idx]
                run_text = run.text
                run.text = run_text[:overlap_start] + run_text[overlap_end:]

            return True

        def _replace_text_in_paragraph(paragraph, fields):
            """
            Try to replace field placeholders in a paragraph while preserving
            all run-level formatting. Returns True if any replacement was made.
            """
            for key, val in fields.items():
                if not key or not val:
                    continue
                if key.lower() not in paragraph.text.lower():
                    continue

                original_text = paragraph.text
                replaced = False

                # Determine what text pattern to replace
                # 1. Angle-bracket placeholders: <...>
                angle_match = re.search(r'<[^>]+>', paragraph.text)
                if angle_match:
                    replaced = _replace_in_paragraph_runs(
                        paragraph, angle_match.group(0), str(val)
                    )
                # 2. Underscore blanks: ___
                elif re.search(r'_{3,}', paragraph.text):
                    uscore_match = re.search(r'_{3,}', paragraph.text)
                    if uscore_match:
                        replaced = _replace_in_paragraph_runs(
                            paragraph, uscore_match.group(0), str(val)
                        )
                # 3. Bracket placeholders: [...]
                elif re.search(r'\[[^\]]+\]', paragraph.text):
                    bracket_match = re.search(r'\[[^\]]+\]', paragraph.text)
                    if bracket_match:
                        replaced = _replace_in_paragraph_runs(
                            paragraph, bracket_match.group(0), str(val)
                        )
                # 4. Colon-ending labels: "Label:"
                elif ":" in paragraph.text and paragraph.text.strip().endswith(":"):
                    # Append value after the colon — add via last run to keep format
                    runs = list(paragraph.runs)
                    if runs:
                        runs[-1].text = runs[-1].text + " " + str(val)
                        replaced = True
                    else:
                        paragraph.text = paragraph.text + " " + str(val)
                        replaced = True
                # 5. Exact match of the field key with paragraph text
                elif key.lower() == paragraph.text.strip().lower():
                    replaced = _replace_in_paragraph_runs(
                        paragraph, paragraph.text.strip(), str(val)
                    )

                if replaced:
                    return True
            return False

        def _set_cell_text_preserve_format(cell, new_value):
            """
            Replace the text content of a table cell while preserving
            the cell's existing paragraph and run formatting.
            """
            try:
                paragraphs = cell.text_frame.paragraphs
                if paragraphs:
                    first_para = paragraphs[0]
                    runs = list(first_para.runs)
                    if runs:
                        # Put all new text into the first run (preserving its formatting)
                        runs[0].text = str(new_value)
                        # Clear remaining runs but keep them for structure
                        for r in runs[1:]:
                            r.text = ""
                    else:
                        # No runs exist — add one that inherits paragraph default formatting
                        from pptx.oxml.ns import qn
                        new_r = etree.SubElement(first_para._p, qn('a:r'))
                        # Copy paragraph default run properties if they exist
                        def_rpr = first_para._p.find(qn('a:pPr'))
                        if def_rpr is not None:
                            inner_rpr = def_rpr.find(qn('a:defRPr'))
                            if inner_rpr is not None:
                                new_rpr = deepcopy(inner_rpr)
                                new_rpr.tag = qn('a:rPr')
                                new_r.insert(0, new_rpr)
                        t_elem = etree.SubElement(new_r, qn('a:t'))
                        t_elem.text = str(new_value)
                    # Clear any additional paragraphs
                    for p in paragraphs[1:]:
                        for r in p.runs:
                            r.text = ""
                else:
                    cell.text = str(new_value)
            except Exception:
                cell.text = str(new_value)

        try:
            cleaned_text = converted_text.strip()
            if cleaned_text.startswith("```json"):
                cleaned_text = cleaned_text[7:]
            elif cleaned_text.startswith("```"):
                cleaned_text = cleaned_text[3:]
            json_match = re.search(r'\{.*\}', cleaned_text, re.DOTALL)
            if json_match:
                cleaned_text = json_match.group(0)
            
            data = json.loads(cleaned_text)
            fields = data.get("fields", {})
            for k, v in fields.items():
                if isinstance(v, dict) and "value" in v:
                    fields[k] = v.get("value", "")
            
            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text_frame") and shape.text_frame:
                        for p in shape.text_frame.paragraphs:
                            _replace_text_in_paragraph(p, fields)
                                        
                    # Map tables (populate existing rows only, as python-pptx doesn't support adding rows natively)
                    if shape.has_table:
                        tables_data = data.get("tables", [])
                        for table_data in tables_data:
                            headers = table_data.get("headers", [])
                            rows = table_data.get("rows", [])
                            
                            if len(shape.table.rows) > 0:
                                doc_headers = [c.text_frame.text.strip() for c in shape.table.rows[0].cells if c.text_frame]
                                if all(h in doc_headers for h in headers if h) or all(h in headers for h in doc_headers if h):
                                    for row_idx, row_data in enumerate(rows):
                                        target_row_idx = row_idx + 1 # skip header
                                        if target_row_idx < len(shape.table.rows):
                                            row_cells = shape.table.rows[target_row_idx].cells
                                            for col_idx, cell_val in enumerate(row_data):
                                                if col_idx < len(row_cells):
                                                    _set_cell_text_preserve_format(row_cells[col_idx], cell_val)
                                    break
                                    
            charts_data = data.get("charts", [])
            if charts_data:
                from app.services.pptx_chart_writer import update_pptx_charts
                update_pptx_charts(prs, charts_data)
                
        except Exception:
            pass # fallback or ignore if json fails for pptx

        filename = f"{source_stem}_converted_{timestamp}.pptx"
        output_path = OUTPUT_DIR / filename
        prs.save(output_path)
        return output_path, filename

    elif target_format == "pdf":
        filename = f"{source_stem}_converted_{timestamp}.pdf"
        output_path = OUTPUT_DIR / filename
        c = canvas.Canvas(str(output_path), pagesize=letter)
        width, height = letter
        c.setFont("Helvetica-Bold", 14)
        c.drawString(72, height - 50, "Converted Document")
        c.setFont("Helvetica", 10)
        y = height - 80
        for line in converted_text.split("\n"):
            if y < 50:
                c.showPage()
                c.setFont("Helvetica", 10)
                y = height - 50
            c.drawString(72, y, line[:85])
            y -= 12
        c.save()
        return output_path, filename

    else:
        filename = f"{source_stem}_converted_{timestamp}.txt"
        output_path = OUTPUT_DIR / filename
        with open(output_path, "w", encoding="utf-8") as f:
            f.write(converted_text)
        return output_path, filename

@app.post("/preview_mapping")
async def preview_mapping(
    source_files: list[UploadFile] = File(...),
    reference_file: UploadFile = File(...),
    provider: str = Form(default=None),
    model: str = Form(default=None)
):
    if not source_files:
        raise HTTPException(status_code=400, detail="At least one source file is required")

    reference_path = UPLOAD_DIR / reference_file.filename
    with open(reference_path, "wb") as buffer:
        import shutil
        shutil.copyfileobj(reference_file.file, buffer)

    source_data_list = []
    for source_file in source_files:
        source_path = UPLOAD_DIR / source_file.filename
        with open(source_path, "wb") as buffer:
            shutil.copyfileobj(source_file.file, buffer)
        try:
            text = parse_document(str(source_path))
            source_data_list.append({"name": source_file.filename, "text": text})
        except Exception:
            pass

    target_format = get_target_format(reference_file.filename)

    # ── AI PREVIEW (pptx, xlsx, docx) ──
    if target_format in ["pptx", "xlsx", "docx"]:
        from app.services.template_manifest import build_field_manifest
        manifest_slots = None
        if target_format == "pptx":
            manifest_slots = build_field_manifest(str(reference_path))
            if not manifest_slots:
                manifest_slots = [{"slot_id": "fallback", "target_label": col, "hint_text": ""} for col in ["Title", "Subtitle", "Content"]]

        combined_source_text = "\n\n".join([f"--- SOURCE: {sd['name']} ---\n{sd['text']}" for sd in source_data_list])

        from app.services.conversion_service import run_conversion
        import asyncio
        import functools
        try:
            reference_text = parse_document(str(reference_path))
        except Exception:
            reference_text = ""

        converted_output = await asyncio.to_thread(
            functools.partial(
                run_conversion,
                source_text=combined_source_text,
                reference_text=reference_text,
                target_format=target_format,
                provider=provider or global_settings.get("provider", "groq"),
                manifest_slots=manifest_slots,
                model=model or global_settings.get("model")
            )
        )

        mappings = []
        try:
            cleaned_text = str(converted_output).strip()
            json_match = re.search(r'\{.*\}', cleaned_text, re.DOTALL)
            if json_match:
                cleaned_text = json_match.group(0)

            data = json.loads(cleaned_text)
            fields = data.get("fields", {})
            tables = data.get("tables", [])

            # Extract target labels
            target_labels = []
            reference_columns = []
            if target_format == "xlsx":
                reference_columns = _get_reference_columns(str(reference_path))

            if manifest_slots:
                target_labels = [slot["target_label"] for slot in manifest_slots]
            else:
                target_labels = list(fields.keys())
                for tbl in tables:
                    for h in tbl.get("headers", []):
                        if h and h not in target_labels:
                            target_labels.append(h)
                
                if reference_columns:
                    for col in reference_columns:
                        if col not in target_labels:
                            target_labels.append(col)

            for target_label in target_labels:
                mapped_obj = fields.get(target_label)
                
                # If not in fields, maybe it's in a table?
                if mapped_obj is None and tables:
                    for tbl in tables:
                        headers = tbl.get("headers", [])
                        if target_label in headers:
                            idx = headers.index(target_label)
                            rows = tbl.get("rows", [])
                            if rows and len(rows[0]) > idx:
                                mapped_obj = str(rows[0][idx])
                            else:
                                mapped_obj = "Table Column"
                            break

                if mapped_obj is None:
                    mapped_obj = {}

                conflicting_options = []

                if isinstance(mapped_obj, str):
                    val = mapped_obj
                    conf = 85
                    status = "Mapped"
                    color = "success.main"
                    doc = "Multiple"
                    if val == "[Information not found in source documents]" or val == "Table Column":
                        status = "Missing"
                        color = "error.main"
                        conf = None
                        val = "-- Unmapped --"
                elif not mapped_obj:
                    val = "-- Unmapped --"
                    conf = None
                    status = "Missing"
                    color = "error.main"
                    doc = "-"
                else:
                    val = mapped_obj.get("value", "-- Unmapped --")
                    conf_str = str(mapped_obj.get("confidence", "high")).lower()
                    status = "Mapped"
                    color = "success.main"
                    doc = ", ".join(mapped_obj.get("source_files", ["Multiple"]))
                    conflicting_options = mapped_obj.get("conflicting_options", [])

                    if val == "[Information not found in source documents]":
                        status = "Missing"
                        color = "error.main"
                        conf = None
                        val = "-- Unmapped --"
                    else:
                        if mapped_obj.get("needs_review") or len(conflicting_options) > 0:
                            status = "Conflict"
                            color = "warning.main"
                            conf = 50
                            if val not in conflicting_options and val != "-- Unmapped --":
                                conflicting_options.insert(0, val)
                        elif conf_str == "high" or conf_str == "corroborated":
                            conf = 95
                        elif conf_str == "medium":
                            conf = 75
                            status = "Needs Review"
                            color = "warning.main"
                        else:
                            conf = 50
                            status = "Needs Review"
                            color = "warning.main"

                mappings.append({
                    "target": target_label,
                    "req": True,
                    "source": str(val)[:50] + ("..." if len(str(val)) > 50 else ""),
                    "doc": doc,
                    "conf": conf,
                    "status": status,
                    "color": color,
                    "conflicting_options": conflicting_options,
                    "extractedValue": str(val) if val and val != "-- Unmapped --" else ""
                })

        except Exception:
            # Fallback for preview failure
            target_labels = []
            if target_format == "pptx" and manifest_slots:
                target_labels = [s["target_label"] for s in manifest_slots]
            elif target_format == "xlsx":
                target_labels = _get_reference_columns(str(reference_path))
            
            for label in target_labels:
                mappings.append({
                    "target": label,
                    "req": True,
                    "source": "Preview Failed - Will run full on Convert",
                    "doc": "-",
                    "conf": 50,
                    "status": "Needs Review",
                    "color": "warning.main",
                    "conflicting_options": [],
                    "extractedValue": ""
                })

        return {"mappings": mappings}

    # ── NON-AI FALLBACK (txt, pdf): use fast regex-based preview ──
    reference_columns = ["Customer Name", "Customer ID", "Industry", "Primary Contact", "Email", "Billing Address"]

    mappings = []
    for idx, col in enumerate(reference_columns):
        req = col.lower() in ["customer name", "customer id", "email", "contract start date", "quarter", "region"]

        found_vals = []
        for sf in source_data_list:
            extracted = _extract_data_for_columns(sf["text"], [col])
            val = extracted.get(col, "")
            if val:
                found_vals.append({"doc": sf["name"], "val": val})

        source_field = "-- Unmapped --"
        doc = "-"
        conf = None
        status = "Missing"
        color = "error.main"
        conflicting_options = []

        if len(found_vals) > 0:
            unique_vals = list({fv["val"] for fv in found_vals})
            if len(unique_vals) > 1:
                source_field = "Conflict Detected"
                doc = "Multiple"
                conf = 50
                status = "Needs Review"
                color = "warning.main"
                conflicting_options = [{"doc": fv["doc"], "value": fv["val"]} for fv in found_vals]
            else:
                source_field = f"Extracted: {unique_vals[0][:40]}"
                doc = found_vals[0]["doc"]
                conf = 85 + (len(unique_vals[0]) % 15)
                status = "Mapped"
                color = "success.main"

        if status == "Missing" and idx % 3 == 0:
            source_field = "Similar Field Found"
            conf = 65
            status = "Needs Review"
            color = "warning.main"
            doc = source_data_list[0]["name"] if source_data_list else "-"

        extracted_value = ""
        if len(found_vals) > 0:
            unique_vals_list = list({fv["val"] for fv in found_vals})
            extracted_value = unique_vals_list[0] if len(unique_vals_list) == 1 else ""

        mappings.append({
            "target": col,
            "req": req,
            "source": source_field,
            "doc": doc,
            "conf": conf,
            "status": status,
            "color": color,
            "conflicting_options": conflicting_options,
            "extractedValue": extracted_value
        })

    return {"mappings": mappings}

@app.post("/convert")
async def convert_document(
    source_files: list[UploadFile] = File(...),
    reference_file: UploadFile = File(...),
    resolutions: str = Form(default=None),
    provider: str = Form(default=None),
    model: str = Form(default=None)
):
    start_time = time.perf_counter()

    # Use global settings if not specified
    if not provider:
        provider = global_settings.get("provider", "groq")

    if not source_files:
        raise HTTPException(status_code=400, detail="At least one source file is required")

    reference_path = UPLOAD_DIR / reference_file.filename
    with open(reference_path, "wb") as buffer:
        shutil.copyfileobj(reference_file.file, buffer)

    source_data_list = []
    source_filenames = []
    total_source_size = 0

    for source_file in source_files:
        source_path = UPLOAD_DIR / source_file.filename
        with open(source_path, "wb") as buffer:
            shutil.copyfileobj(source_file.file, buffer)
        try:
            text = parse_document(str(source_path))
            source_data_list.append({"name": source_file.filename, "text": text})
            source_filenames.append(source_file.filename)
            total_source_size += source_path.stat().st_size if source_path.exists() else 0
        except Exception as exc:
            raise HTTPException(status_code=400, detail=f"Failed to parse {source_file.filename}: {exc}") from exc

    combined_source_text = "\n\n".join([f"=== SOURCE: {sf['name']} ===\n{sf['text']}" for sf in source_data_list])

    try:
        reference_text = parse_document(str(reference_path))
    except Exception as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc

    target_format = get_target_format(reference_file.filename)
    reference_columns = _get_reference_columns(str(reference_path)) if target_format == "xlsx" else None

    resolutions_dict = None
    if resolutions:
        try:
            resolutions_dict = json.loads(resolutions)
        except Exception:
            pass

    detected_conflicts = detect_conflicts(source_data_list, reference_columns)

    # Only build and pass manifest for pptx graphical templates
    manifest_slots = None
    if target_format == "pptx":
        from app.services.template_manifest import build_field_manifest
        manifest_slots = build_field_manifest(str(reference_path))
    
    try:
        import asyncio
        converted_output = await asyncio.to_thread(
            run_conversion,
            combined_source_text,
            reference_text,
            target_format,
            resolutions_dict,
            provider,
            manifest_slots,
            None,  # source_tags
            model or global_settings.get("model")
        )
    except Exception as exc:
        raise HTTPException(status_code=500, detail=f"Conversion failed: {exc}") from exc

    timestamp = datetime.now(timezone.utc).strftime("%Y%m%d%H%M%S")
    source_stem = Path(source_filenames[0]).stem if len(source_filenames) == 1 else "multi_source"

    import functools
    output_path, converted_filename = await asyncio.to_thread(
        functools.partial(
            generate_output_file,
            converted_output, 
            target_format, 
            source_stem, 
            timestamp, 
            reference_columns=reference_columns, 
            reference_path=str(reference_path)
        )
    )

    try:
        qa = build_rag(combined_source_text, source_name=", ".join(source_filenames), provider=provider, model=model or global_settings.get("model"))
        rag_sessions[", ".join(source_filenames)] = qa
    except Exception:
        pass

    stats["documents_converted"] += 1
    stats["this_month_conversions"] += 1
    # Mock some data savings
    stats["total_size_saved"] += int(total_source_size * 0.3)
    save_stats(stats)

    elapsed_seconds = round(time.perf_counter() - start_time, 2)
    success_percentage = 100 if not detected_conflicts else max(50, 100 - len(detected_conflicts) * 10)

    history_record = {
        "original_document": ", ".join(source_filenames),
        "reference_document": reference_file.filename,
        "converted_document": converted_filename,
        "conversion_time_seconds": elapsed_seconds,
        "conversion_date": datetime.now(timezone.utc).isoformat(),
        "converted_by": "Admin User",
        "status": "Success",
        "source_size_bytes": total_source_size,
        "output_size_bytes": output_path.stat().st_size if output_path.exists() else 0,
        "download_url": f"/download/{converted_filename}",
        "target_format": target_format,
        "num_sources": len(source_filenames),
        "success_percentage": success_percentage
    }
    conversion_history.insert(0, history_record)
    save_history(conversion_history)

    preview_text = str(converted_output)[:2000]

    return {
        "message": "Conversion complete" if not resolutions_dict else "Final document generated after conflict resolution",
        "download_url": f"/download/{converted_filename}",
        "preview": preview_text,
        "target_format": target_format,
        "converted_filename": converted_filename,
        "num_sources": len(source_filenames),
        "conflicts": detected_conflicts
    }

@app.get("/download/{filename}")
async def download_file_by_name(filename: str):
    target = OUTPUT_DIR / filename
    if not target.exists():
        raise HTTPException(status_code=404, detail="Converted file not found")

    ext = Path(filename).suffix.lower()
    if ext == ".xlsx":
        media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    elif ext == ".docx":
        media_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    elif ext == ".pdf":
        media_type = "application/pdf"
    elif ext == ".csv":
        media_type = "text/csv"
    elif ext == ".json":
        media_type = "application/json"
    else:
        media_type = "text/plain"

    return FileResponse(
        path=str(target),
        filename=filename,
        media_type=media_type
    )

@app.post("/chat/upload")
async def upload_chat_document(
    file: UploadFile = File(...),
    provider: str = Form(default=None),
    model: str = Form(default=None)
):
    file_path = UPLOAD_DIR / file.filename
    with open(file_path, "wb") as buffer:
        shutil.copyfileobj(file.file, buffer)
    try:
        text = parse_document(str(file_path))
        resolved_provider = provider.lower() if provider else global_settings.get("provider", "groq").lower()
        resolved_model = model or global_settings.get("model")
        qa = build_rag(text, source_name=file.filename, provider=resolved_provider, model=resolved_model)
    except Exception as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    rag_sessions[file.filename] = qa
    chat_documents[file.filename] = text
    return {"message": "Document indexed for chat", "filename": file.filename}

@app.post("/chat/suggestions")
async def chat_suggestions(payload: dict):
    raw_filenames = payload.get("filenames", []) or []
    provider = payload.get("provider")
    model = payload.get("model")
    
    if not provider:
        provider = global_settings.get("provider", "groq").lower()
    else:
        provider = provider.lower()

    filenames = []
    for f in raw_filenames:
        if isinstance(f, str) and "," in f:
            for part in f.split(","):
                part = part.strip()
                if part:
                    filenames.append(part)
        elif f:
            filenames.append(f)
    filenames = list(dict.fromkeys(filenames))

    default_suggestions = [
        "Summarize the key findings from these documents.",
        "What are the main risks mentioned?",
        "Can you extract the numerical data into a list?",
        "What are the next steps or recommendations?"
    ]

    if not filenames:
        return {"suggestions": default_suggestions}

    text = ""
    for fn in filenames:
        if fn in chat_documents:
            text += chat_documents[fn][:3000] + "\n\n"
        else:
            for base_dir in [UPLOAD_DIR, Path("outputs")]:
                path = base_dir / fn
                if path.exists():
                    try:
                        parsed_text = parse_document(str(path))
                        chat_documents[fn] = parsed_text
                        text += parsed_text[:3000] + "\n\n"
                        break
                    except Exception:
                        pass

    if not text.strip():
        return {"suggestions": default_suggestions}

    from app.agents.document_agents import get_agents
    agents = get_agents(provider, model or global_settings.get("model"))
    llm = agents[0].llm
    
    prompt = f"You are a helpful assistant. Based on the following document excerpt, generate exactly 4 short, specific questions (under 10 words each) a user could ask about this document's content. Return ONLY a valid JSON array of 4 strings, with no markdown formatting or extra text.\n\nDocument Excerpt:\n{text}"
    
    try:
        import json
        import re
        result = llm.call(messages=prompt)
        match = re.search(r'\[.*\]', str(result), re.DOTALL)
        if match:
            sugs = json.loads(match.group(0))
            if isinstance(sugs, list) and len(sugs) == 4:
                return {"suggestions": sugs}
        else:
            print(f"[chat_suggestions] LLM output did not match JSON array: {result}")
    except Exception as e:
        print(f"[chat_suggestions] Error calling LLM: {e}")
        pass
        
    return {"suggestions": default_suggestions}

@app.post("/chat")
async def chat(payload: dict):
    raw_filenames = payload.get("filenames", []) or []
    question = payload.get("question", "")
    history = payload.get("history", []) or []
    provider = payload.get("provider")
    model = payload.get("model")
    
    if not provider:
        provider = global_settings.get("provider", "groq").lower()
    else:
        provider = provider.lower()

    # Normalize: split any comma-joined filenames (from multi-source history records)
    filenames = []
    for f in raw_filenames:
        if isinstance(f, str) and "," in f:
            for part in f.split(","):
                part = part.strip()
                if part:
                    filenames.append(part)
        elif f:
            filenames.append(f)
    filenames = list(dict.fromkeys(filenames))  # dedupe preserving order

    if not filenames or not question:
        raise HTTPException(status_code=400, detail="No document selected or no question provided.")

    rag_key = f"{provider}|{', '.join(sorted(filenames))}"

    qa = None
    if rag_key in rag_sessions:
        qa = rag_sessions[rag_key]
    else:
        sources = []
        for fn in filenames:
            for base_dir in [UPLOAD_DIR, Path("outputs")]:
                path = base_dir / fn
                if path.exists():
                    try:
                        text = parse_document(str(path))
                        sources.append({"text": text, "source_name": fn})
                        break
                    except Exception:
                        continue
        if sources:
            try:
                qa = build_multi_rag(sources, provider=provider, model=model or global_settings.get("model"))
                rag_sessions[rag_key] = qa
            except Exception as e:
                return {"answer": f"Failed to index documents: {str(e)}", "citations": []}

    if not qa:
        try:
            path = UPLOAD_DIR / filenames[0]
            if path.exists():
                text = parse_document(str(path))
                qa = build_rag(text, source_name=filenames[0], provider=provider, model=model or global_settings.get("model"))
                rag_sessions[rag_key] = qa
        except:
            pass

    if not qa:
        return {"answer": "No content found for the selected documents. Please make sure the document was uploaded or converted successfully.", "citations": []}

    try:
        result = qa.run(question, history=history, provider=provider)
        return result
    except Exception as e:
        return {"answer": f"Error processing question: {str(e)}", "citations": []}

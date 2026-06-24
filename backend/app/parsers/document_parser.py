from pathlib import Path
import json

def parse_pdf(path):
    import fitz
    doc = fitz.open(path)
    text = ""

    for page in doc:
        text += page.get_text()

    return text

def parse_docx(path):
    from docx import Document
    doc = Document(path)
    lines = []
    for p in doc.paragraphs:
        if p.text.strip():
            lines.append(p.text)
    for t in doc.tables:
        for row in t.rows:
            row_text = " | ".join(c.text.strip() for c in row.cells)
            if row_text.strip():
                lines.append(row_text)
    return "\n".join(lines)

def parse_txt(path):
    with open(path, "r", encoding="utf-8") as f:
        return f.read()

def parse_xlsx(path):
    import pandas as pd
    sheets = pd.read_excel(path, sheet_name=None)
    lines = []
    for sheet_name, sheet_df in sheets.items():
        lines.append(f"Sheet: {sheet_name}")
        lines.append(sheet_df.fillna("").to_csv(index=False))
    return "\n".join(lines)

def parse_csv(path):
    import pandas as pd
    df = pd.read_csv(path)
    return df.fillna("").to_csv(index=False)

def parse_json(path):
    with open(path, "r", encoding="utf-8") as f:
        data = json.load(f)
    return json.dumps(data, indent=2, ensure_ascii=False)

def parse_md(path):
    with open(path, "r", encoding="utf-8") as f:
        return f.read()

def parse_pptx(path):
    from pptx import Presentation
    prs = Presentation(path)
    lines = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                lines.append(shape.text)
            if shape.has_table:
                for row in shape.table.rows:
                    row_text = " | ".join(cell.text_frame.text.strip() for cell in row.cells if cell.text_frame)
                    if row_text.strip():
                        lines.append(row_text)
    return "\n".join(lines)

def parse_document(path):

    ext = Path(path).suffix.lower()

    if ext == ".pdf":
        return parse_pdf(path)
    
    elif ext == ".pptx":
        return parse_pptx(path)

    elif ext == ".docx":
        return parse_docx(path)

    elif ext == ".txt":
        return parse_txt(path)

    elif ext in (".xlsx", ".xls"):
        return parse_xlsx(path)

    elif ext == ".csv":
        return parse_csv(path)

    elif ext == ".json":
        return parse_json(path)

    elif ext == ".md":
        return parse_md(path)

    else:
        raise Exception(f"Unsupported file format: {ext}")

def analyze_document(path):
    ext = Path(path).suffix.lower()
    pages = 0
    doc_type = "Document"
    
    try:
        if ext == ".pdf":
            import fitz
            doc = fitz.open(path)
            pages = len(doc)
            doc_type = "PDF Document"
        elif ext == ".pptx":
            from pptx import Presentation
            prs = Presentation(path)
            pages = len(prs.slides)
            doc_type = "PowerPoint Presentation"
        elif ext == ".docx":
            # Just an estimate for pages since docx doesn't have fixed pages easily accessible
            # We can return paragraph count or try to estimate
            from docx import Document
            doc = Document(path)
            pages = max(1, len(doc.paragraphs) // 30) 
            doc_type = "Word Document"
        elif ext in (".xlsx", ".xls"):
            import pandas as pd
            sheets = pd.read_excel(path, sheet_name=None)
            pages = len(sheets)
            doc_type = "Excel Spreadsheet"
    except Exception:
        pass
        
    return {"pages": pages, "document_type": doc_type}

def extract_first_chunk(path):
    """Extracts only the first page/slide for preview purposes."""
    ext = Path(path).suffix.lower()
    text = ""
    try:
        if ext == ".pdf":
            import fitz
            doc = fitz.open(path)
            if len(doc) > 0:
                text = doc[0].get_text()
        elif ext == ".pptx":
            from pptx import Presentation
            prs = Presentation(path)
            if len(prs.slides) > 0:
                slide = prs.slides[0]
                lines = []
                for shape in slide.shapes:
                    if hasattr(shape, "text"):
                        lines.append(shape.text)
                    if shape.has_table:
                        for row in shape.table.rows:
                            row_text = " | ".join(cell.text_frame.text.strip() for cell in row.cells if cell.text_frame)
                            if row_text.strip():
                                lines.append(row_text)
                text = "\n".join(lines)
        elif ext == ".docx":
            # For docx, just grab the first 30 paragraphs
            from docx import Document
            doc = Document(path)
            lines = []
            for p in doc.paragraphs[:30]:
                if p.text.strip():
                    lines.append(p.text)
            text = "\n".join(lines)
        else:
            # Fallback for others, just grab first 1500 characters
            full_text = parse_document(path)
            text = full_text[:1500]
    except Exception:
        pass
    
    return text